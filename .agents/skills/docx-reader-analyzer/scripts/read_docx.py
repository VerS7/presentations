"""
read_docx.py - Professional Local Word Document Reader & Analyzer
Extracts text, headings, tables, metadata, and converts DOCX to Markdown.
Runs 100% offline using python-docx. Zero external APIs.
"""

import os
import sys
import argparse
from typing import List, Dict, Any
from docx import Document
from docx.opc.exceptions import PackageNotFoundError

def get_docx_metadata(doc: Document) -> Dict[str, Any]:
    """Extracts document core properties."""
    core = doc.core_properties
    return {
        "title": core.title or "(No Title)",
        "author": core.author or "(No Author)",
        "subject": core.subject or "",
        "created": str(core.created) if core.created else "",
        "modified": str(core.modified) if core.modified else "",
        "revision": core.revision or 0,
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables)
    }

def get_docx_outline(doc: Document) -> List[Dict[str, Any]]:
    """Extracts heading hierarchy to understand document structure."""
    outline = []
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if not text:
            continue
        style_name = p.style.name.lower() if p.style else ""
        level = None
        if "heading 1" in style_name:
            level = 1
        elif "heading 2" in style_name:
            level = 2
        elif "heading 3" in style_name:
            level = 3
        elif "title" in style_name:
            level = 0

        if level is not None:
            outline.append({
                "level": level,
                "style": p.style.name,
                "text": text,
                "para_index": i
            })
    return outline

def table_to_markdown(table) -> str:
    """Converts a python-docx Table object to a Markdown table."""
    if not table.rows:
        return ""
    rows_data = []
    for row in table.rows:
        cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
        rows_data.append(cells)

    if not rows_data:
        return ""

    col_count = len(rows_data[0])
    md_lines = []

    # Header
    header = rows_data[0]
    md_lines.append("| " + " | ".join(header) + " |")
    md_lines.append("| " + " | ".join(["---"] * col_count) + " |")

    # Data rows
    for r in rows_data[1:]:
        md_lines.append("| " + " | ".join(r) + " |")

    return "\n".join(md_lines)

def docx_to_markdown(docx_path: str) -> str:
    """Converts an entire DOCX file into structured Markdown."""
    doc = Document(docx_path)
    output_parts = []

    # Map paragraph blocks and tables by their appearance in document body
    # Using document elements iteration
    for child in doc.element.body:
        tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
        if tag == "p":
            # Paragraph
            text_runs = []
            for node in child.iter():
                ntag = node.tag.split("}")[-1] if "}" in node.tag else node.tag
                if ntag == "t" and node.text:
                    text_runs.append(node.text)
            p_text = "".join(text_runs).strip()
            if not p_text:
                continue

            # Detect heading from style
            style_elem = child.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pStyle')
            val = style_elem.attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', '') if style_elem is not None else ''
            val_lower = val.lower()

            if 'title' in val_lower:
                output_parts.append(f"# {p_text}\n")
            elif 'heading1' in val_lower or val == '1':
                output_parts.append(f"## {p_text}\n")
            elif 'heading2' in val_lower or val == '2':
                output_parts.append(f"### {p_text}\n")
            elif 'heading3' in val_lower or val == '3':
                output_parts.append(f"#### {p_text}\n")
            elif 'list' in val_lower or 'bullet' in val_lower:
                output_parts.append(f"- {p_text}")
            else:
                output_parts.append(f"{p_text}\n")

        elif tag == "tbl":
            # Find corresponding table in doc.tables
            for tbl in doc.tables:
                if tbl._element == child:
                    md_tbl = table_to_markdown(tbl)
                    if md_tbl:
                        output_parts.append(f"\n{md_tbl}\n")
                    break

    return "\n".join(output_parts)

def analyze_docx(docx_path: str):
    """Outputs an analytical summary of the document for an AI or user."""
    if not os.path.exists(docx_path):
        print(f"[ERROR] File not found: {docx_path}")
        return

    doc = Document(docx_path)
    meta = get_docx_metadata(doc)
    outline = get_docx_outline(doc)

    all_text = " ".join([p.text.strip() for p in doc.paragraphs if p.text.strip()])
    words = all_text.split()
    word_count = len(words)

    print("=" * 60)
    print(f"DOCX ANALYSIS REPORT: {os.path.basename(docx_path)}")
    print("=" * 60)
    print(f"Title:       {meta['title']}")
    print(f"Author:      {meta['author']}")
    print(f"Created:     {meta['created']}")
    print(f"Word Count:  {word_count} words (~{word_count // 250 + 1} min read)")
    print(f"Paragraphs:  {meta['paragraph_count']}")
    print(f"Tables:      {meta['table_count']}")
    print("-" * 60)
    print("DOCUMENT OUTLINE / HEADINGS:")
    if outline:
        for item in outline:
            indent = "  " * item["level"]
            print(f"{indent}- [{item['style']}] {item['text']}")
    else:
        print("  (No standard heading styles found; plain text document)")
    print("=" * 60)

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Read and analyze Word (.docx) documents")
    parser.add_argument("file", help="Path to .docx file")
    parser.add_argument("--analyze", action="store_true", help="Print summary and heading structure")
    parser.add_argument("--markdown", action="store_true", help="Export full document to Markdown")
    parser.add_argument("--tables", action="store_true", help="Print only extracted tables as Markdown")
    args = parser.parse_args()

    if not os.path.exists(args.file):
        print(f"[ERROR] File not found: {args.file}")
        sys.exit(1)

    if args.markdown:
        print(docx_to_markdown(args.file))
    elif args.tables:
        doc = Document(args.file)
        for i, tbl in enumerate(doc.tables, 1):
            print(f"### Table {i}\n")
            print(table_to_markdown(tbl))
            print("\n")
    else:
        analyze_docx(args.file)
