"""
audit_presentation.py - Pre-flight QA Inspector for PowerPoint Decks
Scans any .pptx file and audits slide structure, word count, density, and speaker notes.
"""

import sys
import os
from pptx import Presentation

def audit_deck(pptx_path: str):
    if not os.path.exists(pptx_path):
        print(f"[ERROR] File not found: {pptx_path}")
        sys.exit(1)

    prs = Presentation(pptx_path)
    slide_count = len(prs.slides)
    width_in = prs.slide_width.inches
    height_in = prs.slide_height.inches

    is_16_9 = abs((width_in / height_in) - (16.0 / 9.0)) < 0.05

    print("=" * 60)
    print(f"PPTX QA AUDIT REPORT: {os.path.basename(pptx_path)}")
    print("=" * 60)
    print(f"Total Slides: {slide_count}")
    print(f"Dimensions:   {width_in:.2f}\" x {height_in:.2f}\" (Aspect Ratio: {'16:9 Widescreen [PASS]' if is_16_9 else 'NON-STANDARD [WARN]'})")
    print("-" * 60)

    total_words = 0
    slides_with_notes = 0
    warnings = []

    for idx, slide in enumerate(prs.slides, 1):
        # Extract slide text
        text_runs = []
        shape_count = len(slide.shapes)
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    t = p.text.strip()
                    if t:
                        text_runs.append(t)

        all_text = " ".join(text_runs)
        words = all_text.split()
        word_count = len(words)
        total_words += word_count

        # Extract title (usually first text run)
        title_text = text_runs[0] if text_runs else "(No Title Detected)"
        if len(title_text) > 50:
            title_text = title_text[:47] + "..."

        # Check speaker notes
        has_notes = False
        notes_word_count = 0
        if slide.has_notes_slide:
            notes_tf = slide.notes_slide.notes_text_frame
            notes_text = notes_tf.text.strip()
            if notes_text:
                has_notes = True
                slides_with_notes += 1
                notes_word_count = len(notes_text.split())

        # Density check
        density_status = "[PASS]"
        if word_count > 90:
            density_status = "[HIGH DENSITY]"
            warnings.append(f"Slide {idx}: High word count ({word_count} words). Consider splitting or card container.")
        elif word_count == 0 and shape_count == 0:
            density_status = "[EMPTY SLIDE]"
            warnings.append(f"Slide {idx}: Empty slide detected.")

        notes_flag = "[NOTES]" if has_notes else "[NO NOTES]"
        print(f"Slide {idx:2d} | Words: {word_count:3d} | Shapes: {shape_count:2d} | {notes_flag} | {title_text}")

    print("-" * 60)
    print(f"Average Words/Slide:  {total_words / max(slide_count, 1):.1f}")
    print(f"Speaker Notes Cover: {slides_with_notes}/{slide_count} slides ({slides_with_notes*100//max(slide_count, 1)}%)")

    print("-" * 60)
    if warnings:
        print(f"AUDIT WARNINGS ({len(warnings)}):")
        for w in warnings:
            print(f"  ! {w}")
    else:
        print("ALL QUALITY CHECKS PASSED: Deck is structurally clean and ready to present.")
    print("=" * 60)

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python audit_presentation.py <deck.pptx>")
        sys.exit(1)
    audit_deck(sys.argv[1])
